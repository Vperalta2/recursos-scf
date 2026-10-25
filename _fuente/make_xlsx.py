#!/usr/bin/env python3
"""Plantillas Excel con fórmulas para el portal, con la imagen de Solución Contable y Fiscal."""
import json, subprocess, datetime as dt
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.drawing.image import Image
from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, AnchorMarker
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).parent
OUT = ROOT / "site/formatos"; OUT.mkdir(parents=True, exist_ok=True)
LOGO = next(p for p in (ROOT / "assets/logo-192.png", ROOT.parent / "assets/logo-192.png") if p.exists())

# Paleta del portal (site.css)
NAVY, NAVY2, GOLD, GOLD_INK, PAPER, INK, MUTED, LINE, TINT = "1B3562", "24467F", "C9AD69", "85662A", "F8F6F0", "18202E", "5A6273", "E2DDCF", "EEF1F7"
INPUT, URGENT, OK = "FFF6DD", "B42B2B", "2C7A4B"
FONT = "Calibri"

def F(size=11, bold=False, color=INK, italic=False, u=None): return Font(name=FONT, size=size, bold=bold, color=color, italic=italic, underline=u)
def fill(c): return PatternFill("solid", fgColor=c)
thin, gold_s, hair = Side(style="thin", color=LINE), Side(style="medium", color=GOLD), Side(style="thin", color=GOLD)
BOX = Border(top=thin, bottom=thin, left=thin, right=thin)
IN_BOX = Border(top=hair, bottom=hair, left=hair, right=hair)
ROW_LINE = Border(bottom=thin)
PESOS = '"$"#,##0'
UNLOCK = Protection(locked=False)

EMPRESA = "SOLUCION CONTABLE Y FISCAL S.A.S."
CONTACTO = "WhatsApp 300 507 4934  ·  solucontfisc@gmail.com  ·  recursos.solucioncontableyfiscal.com"
DIRECCION = "Vía 40 #73-290, Edificio Mix, Oficina 612, Barranquilla · Atención virtual en toda Colombia"
ENLACES = [
    ("Escríbenos por WhatsApp: 300 507 4934", "https://wa.me/573005074934"),
    ("Agenda una cita", "https://calendar.app.google/NaQe5qqNjgy27SrF7"),
    ("Calendario y calculadoras gratis: recursos.solucioncontableyfiscal.com", "https://recursos.solucioncontableyfiscal.com"),
    ("Correo: solucontfisc@gmail.com", "mailto:solucontfisc@gmail.com"),
]
AVISO = "Información de referencia; no reemplaza la asesoría profesional. Verifica tu caso particular con un contador."

def load_data():
    out = subprocess.check_output(["node", "-e", "const D=require('./src/data.js');console.log(JSON.stringify(D))"], cwd=ROOT)
    return json.loads(out)
D = load_data()
C = D["cal"]

def d(mes, dia):
    y, m = map(int, mes.split("-")); return dt.date(y, m, dia)

# ---------- Piezas comunes ----------
def setup(ws, widths, landscape=False, tab=NAVY, one_page=False):
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = tab
    ws.column_dimensions["A"].width = 2.5
    for col, w in widths.items(): ws.column_dimensions[col].width = w
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 1 if one_page else 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.oddFooter.left.text = "Solución Contable y Fiscal · WhatsApp 300 507 4934"
    ws.oddFooter.left.size = 8
    ws.oddFooter.right.text = "Página &P de &N"; ws.oddFooter.right.size = 8

def band(ws, row, last, color, height=None):
    for c in range(1, last + 1): ws.cell(row=row, column=c).fill = fill(color)
    if height: ws.row_dimensions[row].height = height

def header(ws, last, title, subtitle):
    """Franja azul con logo, título, subtítulo y contacto, cerrada con una regla dorada."""
    L = get_column_letter(last)
    for r, h in ((1, 10), (2, 30), (3, 18), (4, 18), (5, 10)): band(ws, r, last, NAVY, h)
    band(ws, 6, last, GOLD, 4)
    # Logo al extremo derecho de la franja (la sangría del texto no es fiable entre Excel y LibreOffice)
    lw = ws.column_dimensions[L].width or 13
    off = max(int(lw * 7 + 5) - 72, 0)
    img = Image(str(LOGO)); img.width = img.height = 64
    img.anchor = OneCellAnchor(_from=AnchorMarker(col=last - 1, colOff=pixels_to_EMU(off), row=1, rowOff=pixels_to_EMU(2)),
                               ext=XDRPositiveSize2D(pixels_to_EMU(64), pixels_to_EMU(64)))
    ws.add_image(img)
    for r, text, font in ((2, title, F(18, True, "FFFFFF")), (3, subtitle, F(10.5, False, GOLD)), (4, CONTACTO, F(9.5, False, "D7DEEA"))):
        ws.merge_cells(f"B{r}:{L}{r}")
        c = ws[f"B{r}"]; c.value = text; c.font = font
        c.alignment = Alignment(indent=1, vertical="center")

def section(ws, row, last, text, n=None):
    L = get_column_letter(last)
    ws.merge_cells(f"B{row}:{L}{row}")
    c = ws[f"B{row}"]; c.value = (f"{n}.  " if n else "") + text.upper()
    c.font = F(10, True, GOLD_INK); c.alignment = Alignment(vertical="bottom")
    for col in range(2, last + 1): ws.cell(row=row, column=col).border = Border(bottom=gold_s)
    ws.row_dimensions[row].height = 24

def legend(ws, row, last):
    ws[f"B{row}"] = "  Tus datos  "; ws[f"B{row}"].fill = fill(INPUT); ws[f"B{row}"].border = IN_BOX
    ws[f"B{row}"].font = F(9.5, True, NAVY)
    ws[f"B{row}"].alignment = Alignment(horizontal="left")
    ws[f"C{row}"] = "Escribe solo en las celdas amarillas; el resto se calcula solo."
    ws[f"C{row}"].font = F(9.5, False, MUTED, italic=True)

def label(ws, ref, text, bold=False, note=None, note_ref=None):
    ws[ref] = text; ws[ref].font = F(11, bold, NAVY if bold else INK); ws[ref].alignment = Alignment(vertical="center", wrap_text=True)
    if note and note_ref:
        ws[note_ref] = note; ws[note_ref].font = F(9, False, MUTED); ws[note_ref].alignment = Alignment(vertical="center", wrap_text=True, indent=1)

def inp(c, value, fmt=None):
    c.value = value; c.fill = fill(INPUT); c.border = IN_BOX; c.font = F(12, True, NAVY); c.protection = UNLOCK
    c.alignment = Alignment(horizontal="right", vertical="center")
    if fmt: c.number_format = fmt

def out(c, formula, fmt=None, strong=False):
    c.value = formula; c.alignment = Alignment(horizontal="right", vertical="center")
    c.font = F(11, strong, NAVY if strong else INK)
    if fmt: c.number_format = fmt

def result(ws, row, last_label_col, value_col, text, formula, fmt):
    """Tarjeta de resultado: fila azul con cifra grande."""
    for col in range(2, value_col + 1): ws.cell(row=row, column=col).fill = fill(NAVY)
    lc = ws.cell(row=row, column=2, value=text); lc.font = F(13, True, "FFFFFF"); lc.alignment = Alignment(vertical="center", indent=1)
    if last_label_col > 2: ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=last_label_col)
    vc = ws.cell(row=row, column=value_col, value=formula); vc.font = F(15, True, GOLD); vc.number_format = fmt
    vc.alignment = Alignment(horizontal="right", vertical="center")
    ws.row_dimensions[row].height = 34

def footer(ws, row, last, sources=None):
    L = get_column_letter(last)
    if sources:
        ws.merge_cells(f"B{row}:{L}{row}"); c = ws[f"B{row}"]; c.value = sources
        c.font = F(9, False, MUTED, italic=True); c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row].height = 28; row += 1
    row += 1
    for r in range(row, row + 8): band(ws, r, last, PAPER)
    for col in range(1, last + 1): ws.cell(row=row, column=col).border = Border(top=gold_s)
    lines = [("¿Prefieres que lo hagamos por ti?", F(13, True, NAVY), None),
             ("Contabilidad · Impuestos · Revisoría fiscal y auditoría · Nómina electrónica", F(10, False, GOLD_INK, ), None)]
    lines += [(t, F(10.5, False, NAVY2, u="single"), url) for t, url in ENLACES]
    lines += [(f"{EMPRESA} · {DIRECCION}", F(9, False, MUTED), None), (AVISO, F(8.5, False, MUTED, italic=True), None)]
    ws.row_dimensions[row].height = 8; row += 1
    for t, font, url in lines:
        ws.merge_cells(f"B{row}:{L}{row}"); c = ws[f"B{row}"]; c.value = t; c.font = font
        if url: c.hyperlink = url
        band(ws, row, last, PAPER); row += 1
    band(ws, row, last, PAPER, 8)
    return row

def protect(ws):
    ws.protection.sheet = True
    ws.protection.autoFilter = False; ws.protection.sort = False
    ws.protection.formatColumns = False; ws.protection.formatRows = False

def props(wb, title, subject):
    wb.properties.creator = "Solución Contable y Fiscal S.A.S."
    wb.properties.title = title; wb.properties.subject = subject
    wb.properties.keywords = "Solución Contable y Fiscal, Colombia, DIAN, 2026"

# ---------- 1. Calendario por NIT ----------
def calendario():
    wb = Workbook(); ws = wb.active; ws.title = "Mis fechas"
    dd = wb.create_sheet("Datos DIAN 2026")
    props(wb, "Calendario tributario 2026 por NIT", "Fechas DIAN 2026 según los dígitos del NIT")

    # Hoja de datos
    dd.sheet_view.showGridLines = False; dd.sheet_properties.tabColor = GOLD
    dd.append(["Obligación", "Período", "Mes de vencimiento"] + [1, 2, 3, 4, 5, 6, 7, 8, 9, 0])
    rows = []
    def add(obl, per, mes, dias):
        dd.append([obl, per, mes] + [d(mes, x) for x in dias]); rows.append((obl, per))
    for it in C["renta_gc"]: add(it["tipo"], it["detalle"], it["mes"], it["dias"])
    for it in C["renta_pj"]: add(it["tipo"], it["detalle"], it["mes"], it["dias"])
    for it in C["iva_bim"]: add("IVA bimestral", it["periodo"], it["mes"], it["dias"])
    for it in C["iva_cuat"]: add("IVA cuatrimestral", it["periodo"], it["mes"], it["dias"])
    for it in C["retencion"]: add("Retención en la fuente", it["periodo"], it["mes"], it["dias"])
    for it in C["rst_anticipo"]: add("Régimen Simple · anticipo", it["periodo"], it["mes"], it["dias"])
    an = C["rst_anual"]["dias"]; iv = C["rst_iva_anual"]["dias"]
    add("Régimen Simple · declaración anual", "año 2025", C["rst_anual"]["mes"], [an[0], an[0], an[1], an[1], an[2], an[2], an[3], an[3], an[4], an[4]])
    add("Régimen Simple · IVA anual consolidado", "año 2025", C["rst_iva_anual"]["mes"], [iv[0], iv[0], iv[1], iv[1], iv[2], iv[2], iv[3], iv[3], iv[4], iv[4]])
    n = len(rows)
    for r in dd.iter_rows(min_row=2, min_col=4, max_col=13):
        for c in r: c.number_format = "dd/mm/yyyy"
    for r in dd.iter_rows(min_row=2, max_col=13):
        for c in r: c.border = ROW_LINE; c.font = F(10)
    for c in dd[1]: c.font = F(10, True, "FFFFFF"); c.fill = fill(NAVY); c.alignment = Alignment(horizontal="center")
    dd["P1"] = "Dos últimos dígitos"; dd["Q1"] = "Renta personas naturales"
    for ref in ("P1", "Q1"): dd[ref].font = F(10, True, "FFFFFF"); dd[ref].fill = fill(NAVY)
    for k in range(100):
        idx = 49 if k == 0 else (k + 1) // 2 - 1
        y, m, dia = map(int, C["renta_pn"][idx].split("-"))
        dd.cell(row=k + 2, column=16, value=k).number_format = "00"
        dd.cell(row=k + 2, column=17, value=dt.date(y, m, dia)).number_format = "dd/mm/yyyy"
    for col, w in {"A": 38, "B": 30, "C": 18, "P": 18, "Q": 24}.items(): dd.column_dimensions[col].width = w
    for col in "DEFGHIJKLM": dd.column_dimensions[col].width = 12
    dd.freeze_panes = "D2"; protect(dd)
    dd.page_setup.orientation = "landscape"; dd.page_setup.fitToWidth = 1; dd.page_setup.fitToHeight = 0
    dd.sheet_properties.pageSetUpPr.fitToPage = True

    # Hoja principal: columnas B..F
    last = 6
    setup(ws, {"B": 42, "C": 30, "D": 15, "E": 14, "F": 17})
    header(ws, last, "Calendario tributario 2026 por NIT", "Tus fechas DIAN de renta, IVA, retención y Régimen Simple en un solo lugar")
    legend(ws, 8, last)
    section(ws, 10, last, "Tu NIT", 1)
    label(ws, "B11", "Dos últimos dígitos del NIT (sin dígito de verificación)", bold=True)
    inp(ws["C11"], 34, "00"); ws["C11"].font = F(16, True, NAVY); ws.row_dimensions[11].height = 26
    ws["D11"] = "Ejemplo: NIT 900.123.434-5 → 34"; ws["D11"].font = F(9, False, MUTED, italic=True)
    label(ws, "B12", "Último dígito"); out(ws["C12"], "=MOD(C11,10)", "0")
    dv = DataValidation(type="whole", operator="between", formula1="0", formula2="99", showErrorMessage=True,
                        errorTitle="Dígitos del NIT", error="Escribe un número de 00 a 99.", showInputMessage=True,
                        promptTitle="Tu NIT", prompt="Los dos últimos dígitos, sin el dígito de verificación.")
    ws.add_data_validation(dv); dv.add("C11")

    first, lastr = 21, 21 + n - 1
    rngD = f"$D${first}:$D${lastr}"
    section(ws, 14, last, "Resumen", 2)
    label(ws, "B15", "Renta personas naturales (año gravable 2025)", bold=True)
    out(ws["C15"], "=VLOOKUP(MOD(C11,100),'Datos DIAN 2026'!$P$2:$Q$101,2,FALSE)", "dd/mm/yyyy", strong=True)
    label(ws, "B16", "Hoy"); out(ws["C16"], "=TODAY()", "dd/mm/yyyy")
    result(ws, 17, 2, 3, "Tu próximo vencimiento", f'=IFERROR(1/(1/_xlfn.MINIFS({rngD},{rngD},">="&TODAY())),"Sin pendientes")', "dd/mm/yyyy")
    ws.merge_cells("D17:F17")
    c = ws["D17"]; c.value = f'=IFERROR(INDEX($B${first}:$B${lastr},MATCH(C17,{rngD},0))&" · "&INDEX($C${first}:$C${lastr},MATCH(C17,{rngD},0)),"")'
    c.font = F(11, True, "FFFFFF"); c.alignment = Alignment(vertical="center", wrap_text=True, indent=1)
    for col in range(4, last + 1): ws.cell(row=17, column=col).fill = fill(NAVY)
    ws.row_dimensions[17].height = 40

    section(ws, 19, last, "Tus fechas 2026", 3)
    hdr = 20
    for col, t in zip("BCDEF", ["Obligación", "Período", "Fecha límite", "Días que faltan", "Estado"]):
        c = ws[f"{col}{hdr}"]; c.value = t; c.font = F(10.5, True, "FFFFFF"); c.fill = fill(NAVY)
        c.alignment = Alignment(horizontal="left" if col in "BC" else "center", vertical="center")
    ws.row_dimensions[hdr].height = 22
    for i in range(n):
        r = first + i; src = 2 + i
        ws[f"B{r}"] = f"='Datos DIAN 2026'!A{src}"
        ws[f"C{r}"] = f"='Datos DIAN 2026'!B{src}"
        ws[f"D{r}"] = f"=INDEX('Datos DIAN 2026'!$D${src}:$M${src},MATCH($C$12,'Datos DIAN 2026'!$D$1:$M$1,0))"
        ws[f"E{r}"] = f'=IF(D{r}<TODAY(),"",D{r}-TODAY())'
        ws[f"F{r}"] = f'=IF(D{r}<TODAY(),"Vencido",IF(D{r}-TODAY()<=15,"Vence pronto","Pendiente"))'
        ws[f"D{r}"].number_format = "dd/mm/yyyy"; ws[f"E{r}"].number_format = "0"
        for col in "BCDEF":
            c = ws[f"{col}{r}"]; c.border = ROW_LINE; c.font = F(10.5)
            if col in "DEF": c.alignment = Alignment(horizontal="center")
            if i % 2: c.fill = fill(PAPER)
    body = f"B{first}:F{lastr}"
    ws.conditional_formatting.add(body, FormulaRule(formula=[f'$F{first}="Vencido"'], font=Font(color="9AA0AA")))
    ws.conditional_formatting.add(body, FormulaRule(formula=[f'$F{first}="Vence pronto"'], font=Font(color=URGENT, bold=True), fill=fill("FBE6E3")))
    ws.conditional_formatting.add(f"F{first}:F{lastr}", FormulaRule(formula=[f'$F{first}="Pendiente"'], font=Font(color=OK, bold=True)))
    ws.auto_filter.ref = f"B{hdr}:F{lastr}"
    ws.freeze_panes = f"B{first}"
    footer(ws, lastr + 2, last, "Revisa solo las filas que te aplican; usa los filtros del encabezado para ver una obligación. "
           "«Vence pronto» = 15 días o menos. Fuente: Calendario Tributario 2026 de la DIAN (versión del 1 de julio de 2026).")
    ws.print_title_rows = f"{hdr}:{hdr}"
    protect(ws)
    wb.save(OUT / "calendario-tributario-2026-por-nit.xlsx")

# ---------- 2. Retención por salarios ----------
TABLA_383 = [(0, 95, 0, 0, 0), (95, 150, .19, 95, 0), (150, 360, .28, 150, 10), (360, 640, .33, 360, 69), (640, 945, .35, 640, 162), (945, 2300, .37, 945, 268), (2300, 999999999, .39, 2300, 770)]

def retencion():
    wb = Workbook(); ws = wb.active; ws.title = "Retención"
    nom = wb.create_sheet("Varios empleados")
    p = wb.create_sheet("Parámetros")
    props(wb, "Retención en la fuente por salarios 2026", "Procedimiento 1, art. 383 ET")

    # Parámetros
    setup(p, {"B": 24, "C": 16, "D": 16, "E": 16, "F": 16}, tab=GOLD)
    header(p, 6, "Parámetros 2026", "Valores y tabla del art. 383 ET usados en los cálculos")
    label(p, "B8", "UVT 2026", True); out(p["C8"], D["uvt"]["2026"], PESOS, True)
    label(p, "B9", "Salario mínimo 2026", True); out(p["C9"], D["smmlv"], PESOS, True)
    section(p, 11, 6, "Tabla art. 383 ET")
    for col, t in zip("BCDEF", ["Desde (UVT)", "Hasta (UVT)", "Tarifa marginal", "Se resta (UVT)", "Se suma (UVT)"]):
        c = p[f"{col}12"]; c.value = t; c.font = F(10, True, "FFFFFF"); c.fill = fill(NAVY); c.alignment = Alignment(horizontal="center")
    for i, t in enumerate(TABLA_383):
        for j, v in enumerate(t):
            c = p.cell(row=13 + i, column=2 + j, value=v); c.border = ROW_LINE; c.alignment = Alignment(horizontal="center")
            c.number_format = "0%" if j == 2 else "#,##0"
            if i % 2: c.fill = fill(PAPER)
    p["C19"].value = "En adelante"
    footer(p, 21, 6)
    protect(p)
    U, S, TBL = "Parámetros!$C$8", "Parámetros!$C$9", "Parámetros!$B$13:$F$19"

    def ret_uvt(x):
        return f"IF({x}<=95,0,({x}-VLOOKUP({x},{TBL},4,TRUE))*VLOOKUP({x},{TBL},3,TRUE)+VLOOKUP({x},{TBL},5,TRUE))"
    def fsp(x):
        return f"{x}*IF({x}<4*{S},0,IF({x}<16*{S},1%,IF({x}<17*{S},1.2%,IF({x}<18*{S},1.4%,IF({x}<19*{S},1.6%,IF({x}<20*{S},1.8%,2%))))))"

    # Hoja principal
    last = 4
    setup(ws, {"B": 58, "C": 20, "D": 30}, one_page=True)
    header(ws, last, "Retención en la fuente por salarios 2026", "Procedimiento 1 · Arts. 383, 387 y 388 del Estatuto Tributario")
    legend(ws, 8, last)
    section(ws, 10, last, "Datos del empleado", 1)
    inputs = [("salario", "Salario del mes", 15000000, PESOS, ""), ("otros", "Otros pagos laborales del mes", 0, PESOS, "Horas extra, comisiones, bonificaciones"),
              ("dep", "¿Tiene dependientes?", "No", None, "Elige Sí o No · art. 387 ET"),
              ("viv", "Intereses de vivienda (mensual)", 0, PESOS, "Certificado del banco · art. 119 ET"),
              ("prep", "Medicina prepagada (mensual)", 0, PESOS, "Art. 387 ET"),
              ("afc", "Aportes voluntarios a pensión o AFC del mes", 0, PESOS, "Arts. 126-1 y 126-4 ET")]
    R = {}
    for i, (key, text, val, fmt, note) in enumerate(inputs):
        r = 11 + i; R[key] = f"C{r}"; ws.row_dimensions[r].height = 21
        label(ws, f"B{r}", text, note=note, note_ref=f"D{r}"); inp(ws[f"C{r}"], val, fmt)
    dv = DataValidation(type="list", formula1='"Sí,No"', allow_blank=False, showErrorMessage=True, error="Elige Sí o No.")
    ws.add_data_validation(dv); dv.add(R["dep"])
    dvn = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", showErrorMessage=True, error="Escribe un valor en pesos, sin signos.")
    ws.add_data_validation(dvn)
    for k in ("salario", "otros", "viv", "prep", "afc"): dvn.add(R[k])

    section(ws, 18, last, "Depuración de la base", 2)
    calc = [
        ("ing", "Ingresos laborales", "=N({salario})+N({otros})", PESOS, False, ""),
        ("ibc", "Ingreso base de cotización (tope 25 SMMLV)", "=MIN({ing},25*" + S + ")", PESOS, False, ""),
        ("salud", "(−) Salud 4 %", "={ibc}*4%", PESOS, False, "Ingreso no constitutivo · art. 56 ET"),
        ("pens", "(−) Pensión 4 %", "={ibc}*4%", PESOS, False, "Art. 55 ET"),
        ("fsp", "(−) Fondo de solidaridad pensional", "=" + fsp("{ibc}"), PESOS, False, "Desde 4 SMMLV"),
        ("neto", "Ingreso neto", "={ing}-{salud}-{pens}-{fsp}", PESOS, True, ""),
        ("ddep", "(−) Dependientes (10 %, máx. 32 UVT)", '=IF({dep}="Sí",MIN({ing}*10%,32*' + U + "),0)", PESOS, False, ""),
        ("dviv", "(−) Intereses de vivienda (máx. 100 UVT)", "=MIN(N({viv}),100*" + U + ")", PESOS, False, ""),
        ("dprep", "(−) Medicina prepagada (máx. 16 UVT)", "=MIN(N({prep}),16*" + U + ")", PESOS, False, ""),
        ("dafc", "(−) Aportes voluntarios / AFC (máx. 30 % y 3.800 UVT al año)", "=MIN(N({afc}),{ing}*30%,3800*" + U + "/12)", PESOS, False, ""),
        ("sub", "Subtotal", "=MAX({neto}-{ddep}-{dviv}-{dprep}-{dafc},0)", PESOS, False, ""),
        ("ex", "(−) Renta exenta 25 % (máx. 790 UVT al año)", "=MIN({sub}*25%,790*" + U + "/12)", PESOS, False, "Art. 206 num. 10 ET"),
        ("tot", "Total deducciones y rentas exentas", "={ddep}+{dviv}+{dprep}+{dafc}+{ex}", PESOS, False, ""),
        ("lim", "Límite 40 % y 1.340 UVT al año", "=MIN({neto}*40%,1340*" + U + "/12)", PESOS, False, "Art. 336 ET"),
        ("base", "Base gravable", "=MAX({neto}-MIN({tot},{lim}),0)", PESOS, True, ""),
        ("buvt", "Base gravable en UVT", "={base}/" + U, "#,##0.00", False, ""),
        ("tar", "Tarifa marginal", "=VLOOKUP({buvt}," + TBL + ",3,TRUE)", "0%", False, "Tabla art. 383 ET"),
        ("ruvt", "Retención en UVT", "=" + ret_uvt("{buvt}"), "#,##0.00", False, ""),
    ]
    for i, (key, text, f, fmt, strong, note) in enumerate(calc):
        r = 19 + i; R[key] = f"C{r}"
        label(ws, f"B{r}", text, bold=strong, note=note, note_ref=f"D{r}")
        out(ws[f"C{r}"], f.format(**R), fmt, strong)
        for col in "BCD":
            ws[f"{col}{r}"].border = ROW_LINE
            if strong: ws[f"{col}{r}"].fill = fill(TINT)
    r = 19 + len(calc) + 1
    result(ws, r, 2, 3, "Retención del mes", f"={R['ruvt']}*{U}", PESOS)
    ws[f"D{r}"].fill = fill(NAVY)
    footer(ws, r + 2, last, "Topes anuales aplicados en doceavas partes. Para varios empleados usa la hoja «Varios empleados».")
    protect(ws)

    # Varios empleados
    cols = ["Empleado", "Salario del mes", "Otros pagos", "Dependientes", "Intereses vivienda", "Medicina prepagada", "AFC / pensión vol.",
            "Ingresos laborales", "IBC", "Salud y pensión (8 %)", "Solidaridad pensional", "Ingreso neto", "Deducciones",
            "Renta exenta 25 %", "Límite 40 %", "Base gravable", "Base en UVT", "Retención del mes"]
    last = 1 + len(cols)
    setup(nom, {get_column_letter(2 + i): (26 if i == 0 else 15) for i in range(len(cols))}, landscape=True)
    header(nom, last, "Retención por salarios 2026 · varios empleados", "Un empleado por fila: llena las columnas amarillas (B a H)")
    legend(nom, 8, last)
    hdr = 10
    for i, t in enumerate(cols):
        c = nom.cell(row=hdr, column=2 + i, value=t); c.font = F(10, True, "FFFFFF" if i < 7 else NAVY)
        c.fill = fill(NAVY if i < 7 else "D9DFEA"); c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    nom.row_dimensions[hdr].height = 34
    first, n_emp = hdr + 1, 30
    lastr = first + n_emp - 1
    for k in range(n_emp):
        r = first + k
        B, Cc, Dd, E, Fc, G, H, I, J, K, L, M, N, O, P, Q, Rr, Sc = (f"{get_column_letter(2 + i)}{r}" for i in range(18))
        nom[B] = f"Empleado {k + 1}" if k < 2 else None
        nom[Cc] = [3500000, 12000000][k] if k < 2 else None
        nom[E] = "No" if k < 2 else None
        g = lambda f: f'=IF(N({Cc})=0,"",{f})'
        nom[I] = g(f"N({Cc})+N({Dd})")
        nom[J] = g(f"MIN({I},25*{S})")
        nom[K] = g(f"{J}*8%")
        nom[L] = g(fsp(J))
        nom[M] = g(f"{I}-{K}-{L}")
        nom[N] = g(f'IF({E}="Sí",MIN({I}*10%,32*{U}),0)+MIN(N({Fc}),100*{U})+MIN(N({G}),16*{U})+MIN(N({H}),{I}*30%,3800*{U}/12)')
        nom[O] = g(f"MIN(MAX({M}-{N},0)*25%,790*{U}/12)")
        nom[P] = g(f"MIN({M}*40%,1340*{U}/12)")
        nom[Q] = g(f"MAX({M}-MIN({N}+{O},{P}),0)")
        nom[Rr] = g(f"{Q}/{U}")
        nom[Sc] = g(f"{ret_uvt(Rr)}*{U}")
        for i in range(18):
            c = nom.cell(row=r, column=2 + i); c.border = BOX if i < 7 else ROW_LINE; c.font = F(10.5, i == 17, NAVY if i == 17 else INK)
            if i < 7: c.fill = fill(INPUT); c.protection = UNLOCK
            elif k % 2: c.fill = fill(PAPER)
            c.number_format = "#,##0.00" if i == 16 else (PESOS if i not in (0, 3) else "General")
            if i == 3: c.alignment = Alignment(horizontal="center")
    dv = DataValidation(type="list", formula1='"Sí,No"', allow_blank=True, showErrorMessage=True, error="Elige Sí o No.")
    nom.add_data_validation(dv); dv.add(f"E{first}:E{lastr}")
    tr = lastr + 1
    nom[f"B{tr}"] = "Totales"
    for i in range(18):
        c = nom.cell(row=tr, column=2 + i); c.fill = fill(NAVY); c.font = F(10.5, True, GOLD if i == 17 else "FFFFFF")
        if i in (1, 2, 4, 5, 6, 7, 9, 10, 12, 17):
            col = get_column_letter(2 + i); c.value = f"=SUM({col}{first}:{col}{lastr})"; c.number_format = PESOS
    nom.row_dimensions[tr].height = 24
    nom.freeze_panes = f"C{first}"
    nom.print_title_rows = f"{hdr}:{hdr}"
    footer(nom, tr + 2, last, "Mismo cálculo de la hoja «Retención» (procedimiento 1). Deducciones = dependientes + vivienda + prepagada + AFC, con sus topes.")
    protect(nom)
    wb.save(OUT / "retencion-en-la-fuente-salarios-2026.xlsx")

# ---------- 3. Prestaciones sociales ----------
def dias360(a, b):
    return f"DAYS360({a},{b},TRUE)+1+IF(AND(MONTH({b}+1)=3,DAY({b}+1)=1),30-DAY({b}),0)"

def prestaciones():
    wb = Workbook(); ws = wb.active; ws.title = "Prestaciones"
    nom = wb.create_sheet("Varios empleados")
    props(wb, "Liquidador de prestaciones sociales 2026", "Cesantías, intereses, prima y vacaciones")
    last = 4
    setup(ws, {"B": 56, "C": 20, "D": 30}, one_page=True)
    header(ws, last, "Liquidador de prestaciones sociales 2026", "Cesantías, intereses, prima de servicios y vacaciones")
    legend(ws, 8, last)
    section(ws, 10, last, "Datos del trabajador", 1)
    R = {}
    ins = [("sal", "Salario mensual", D["smmlv"], PESOS, "Salario básico"), ("ini", "Fecha de inicio", dt.date(2026, 1, 1), "dd/mm/yyyy", "Ingreso o última liquidación"),
           ("cor", "Fecha de corte", dt.date(2026, 10, 31), "dd/mm/yyyy", "Retiro o fecha a calcular")]
    for i, (k, t, v, fmt, note) in enumerate(ins):
        r = 11 + i; R[k] = f"C{r}"; ws.row_dimensions[r].height = 21
        label(ws, f"B{r}", t, note=note, note_ref=f"D{r}"); inp(ws[f"C{r}"], v, fmt)
    dvd = DataValidation(type="date", operator="between", formula1="DATE(2000,1,1)", formula2="DATE(2030,12,31)", showErrorMessage=True, error="Escribe una fecha válida (dd/mm/aaaa).")
    dvc = DataValidation(type="custom", formula1=f"{R['cor']}>={R['ini']}", showErrorMessage=True, error="La fecha de corte debe ser igual o posterior a la de inicio.")
    ws.add_data_validation(dvd); ws.add_data_validation(dvc); dvd.add(R["ini"]); dvc.add(R["cor"])

    section(ws, 15, last, "Valores 2026", 2)
    label(ws, "B16", "Salario mínimo 2026"); out(ws["C16"], D["smmlv"], PESOS); R["smm"] = "$C$16"
    label(ws, "B17", "Auxilio de transporte 2026"); out(ws["C17"], D["auxTransporte"], PESOS); R["aux"] = "$C$17"

    section(ws, 19, last, "Liquidación", 3)
    calc = [
        ("auxi", "Auxilio de transporte incluido", "=IF({sal}<=2*{smm},{aux},0)", PESOS, "Solo si gana hasta 2 SMMLV"),
        ("base", "Base para cesantías y prima", "={sal}+{auxi}", PESOS, ""),
        ("dias", "Días trabajados", "=" + dias360("{ini}", "{cor}"), "0", "Año de 360 días"),
        ("isem", "Inicio del semestre de la prima", "=MAX({ini},IF(MONTH({cor})<=6,DATE(YEAR({cor}),1,1),DATE(YEAR({cor}),7,1)))", "dd/mm/yyyy", ""),
        ("dsem", "Días de prima del semestre", "=" + dias360("{isem}", "{cor}"), "0", ""),
        ("ces", "Cesantías", "={base}*{dias}/360", PESOS, "Art. 249 CST"),
        ("int", "Intereses sobre cesantías (12 %)", "={ces}*{dias}*12%/360", PESOS, "Ley 52 de 1975"),
        ("pri", "Prima de servicios", "={base}*{dsem}/360", PESOS, "Art. 306 CST"),
        ("vac", "Vacaciones (15 días hábiles por año)", "={sal}*{dias}/720", PESOS, "Art. 186 CST · sin auxilio"),
    ]
    for i, (k, t, f, fmt, note) in enumerate(calc):
        r = 20 + i; R[k] = f"C{r}"
        label(ws, f"B{r}", t, note=note, note_ref=f"D{r}"); out(ws[f"C{r}"], f.format(**R), fmt)
        for col in "BCD": ws[f"{col}{r}"].border = ROW_LINE
    r = 20 + len(calc) + 1
    result(ws, r, 2, 3, "Total prestaciones", f"={R['ces']}+{R['int']}+{R['pri']}+{R['vac']}", PESOS)
    ws[f"D{r}"].fill = fill(NAVY)
    footer(ws, r + 2, last, "Supone que no se han pagado cesantías ni tomado vacaciones en el período. Para varios trabajadores usa la hoja «Varios empleados».")
    protect(ws)

    # Varios empleados
    cols = ["Trabajador", "Salario mensual", "Fecha de inicio", "Fecha de corte", "Auxilio transporte", "Base", "Días trabajados",
            "Inicio semestre", "Días prima", "Cesantías", "Intereses (12 %)", "Prima", "Vacaciones", "Total"]
    last = 1 + len(cols)
    setup(nom, {get_column_letter(2 + i): (26 if i == 0 else 14) for i in range(len(cols))}, landscape=True)
    header(nom, last, "Prestaciones sociales 2026 · varios trabajadores", "Un trabajador por fila: llena las columnas amarillas (B a E)")
    legend(nom, 8, last)
    hdr = 10
    for i, t in enumerate(cols):
        c = nom.cell(row=hdr, column=2 + i, value=t); c.font = F(10, True, "FFFFFF" if i < 4 else NAVY)
        c.fill = fill(NAVY if i < 4 else "D9DFEA"); c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    nom.row_dimensions[hdr].height = 34
    first, n_emp = hdr + 1, 30
    lastr = first + n_emp - 1
    SMM, AUX = "Prestaciones!$C$16", "Prestaciones!$C$17"
    for k in range(n_emp):
        r = first + k
        B, Cc, Dd, E, Fc, G, H, I, J, K, L, M, N, O = (f"{get_column_letter(2 + i)}{r}" for i in range(14))
        if k < 2:
            nom[B] = f"Trabajador {k + 1}"; nom[Cc] = [D["smmlv"], 4200000][k]
            nom[Dd] = [dt.date(2026, 1, 1), dt.date(2026, 3, 16)][k]; nom[E] = dt.date(2026, 10, 31)
        g = lambda f: f'=IF(OR(N({Cc})=0,N({Dd})=0,N({E})=0),"",{f})'
        nom[Fc] = g(f"IF({Cc}<=2*{SMM},{AUX},0)")
        nom[G] = g(f"{Cc}+{Fc}")
        nom[H] = g(dias360(Dd, E))
        nom[I] = g(f"MAX({Dd},IF(MONTH({E})<=6,DATE(YEAR({E}),1,1),DATE(YEAR({E}),7,1)))")
        nom[J] = g(dias360(I, E))
        nom[K] = g(f"{G}*{H}/360")
        nom[L] = g(f"{K}*{H}*12%/360")
        nom[M] = g(f"{G}*{J}/360")
        nom[N] = g(f"{Cc}*{H}/720")
        nom[O] = g(f"{K}+{L}+{M}+{N}")
        for i in range(14):
            c = nom.cell(row=r, column=2 + i); c.border = BOX if i < 4 else ROW_LINE; c.font = F(10.5, i == 13, NAVY if i == 13 else INK)
            if i < 4: c.fill = fill(INPUT); c.protection = UNLOCK
            elif k % 2: c.fill = fill(PAPER)
            c.number_format = "dd/mm/yyyy" if i in (2, 3, 7) else ("0" if i in (6, 8) else (PESOS if i else "General"))
            if i in (2, 3, 6, 7, 8): c.alignment = Alignment(horizontal="center")
    dvd = DataValidation(type="date", operator="between", formula1="DATE(2000,1,1)", formula2="DATE(2030,12,31)", allow_blank=True, showErrorMessage=True, error="Escribe una fecha válida (dd/mm/aaaa).")
    nom.add_data_validation(dvd); dvd.add(f"D{first}:E{lastr}")
    tr = lastr + 1
    nom[f"B{tr}"] = "Totales"
    for i in range(14):
        c = nom.cell(row=tr, column=2 + i); c.fill = fill(NAVY); c.font = F(10.5, True, GOLD if i == 13 else "FFFFFF")
        if i in (1, 9, 10, 11, 12, 13):
            col = get_column_letter(2 + i); c.value = f"=SUM({col}{first}:{col}{lastr})"; c.number_format = PESOS
    nom.row_dimensions[tr].height = 24
    nom.freeze_panes = f"C{first}"
    nom.print_title_rows = f"{hdr}:{hdr}"
    footer(nom, tr + 2, last, "Mismo cálculo de la hoja «Prestaciones». Salario mínimo y auxilio de transporte tomados de esa hoja.")
    protect(nom)
    wb.save(OUT / "liquidador-prestaciones-sociales-2026.xlsx")

calendario(); retencion(); prestaciones()
print(sorted(p.name for p in OUT.iterdir()))
